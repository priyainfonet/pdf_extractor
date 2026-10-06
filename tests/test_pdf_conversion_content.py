from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from app.services import excel_converter, pdf_service, powerpoint_converter, word_converter


OCR_TEXTS = (
    "Invoice",
    "Item",
    "Qty",
    "Price",
    "Apple",
    "2",
    "$3",
    "Orange",
    "4",
    "$6",
)
OCR_BOXES = (
    [[50, 10], [170, 10], [170, 50], [50, 50]],
    [[43, 57], [122, 59], [121, 102], [42, 100]],
    [[555, 54], [620, 62], [615, 107], [550, 99]],
    [[893, 55], [984, 58], [983, 102], [892, 99]],
    [[42, 157], [145, 157], [145, 207], [42, 207]],
    [[554, 160], [584, 160], [584, 198], [554, 198]],
    [[890, 156], [945, 156], [945, 202], [890, 202]],
    [[46, 257], [169, 262], [167, 306], [44, 300]],
    [[555, 262], [584, 262], [584, 296], [555, 296]],
    [[891, 256], [945, 256], [945, 302], [891, 302]],
)


class FakeImage:
    width = 1200
    height = 500

    def convert(self, mode):
        return self


class FakePage:
    images = [{}]
    width = 1200
    height = 500

    def extract_tables(self, table_settings):
        return []

    def extract_text(self):
        return ""

    def to_image(self, resolution):
        return SimpleNamespace(original=FakeImage())


class FakePdf:
    pages = [FakePage()]

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_value, traceback):
        return False


def _blank_graphics_backgrounds(pdf_path):
    image = Image.new("RGB", (1200, 500), "white")
    output = BytesIO()
    image.save(output, format="PNG")
    return [output.getvalue()]


def test_scanned_content_is_preserved_and_excel_extracts_table(tmp_path, monkeypatch):
    monkeypatch.setattr(pdf_service.pdfplumber, "open", lambda _: FakePdf())
    monkeypatch.setattr(pdf_service, "_extract_graphic_backgrounds", _blank_graphics_backgrounds)
    monkeypatch.setattr(
        pdf_service,
        "_get_ocr_engine",
        lambda: lambda image: SimpleNamespace(txts=OCR_TEXTS, boxes=OCR_BOXES),
    )
    page_image = _blank_graphics_backgrounds(None)[0]
    monkeypatch.setattr(word_converter, "render_pdf_pages", lambda _: [page_image])
    monkeypatch.setattr(powerpoint_converter, "render_pdf_pages", lambda _: [page_image])

    pdf_path = str(tmp_path / "scanned.pdf")
    word_path = tmp_path / "converted.docx"
    powerpoint_path = tmp_path / "converted.pptx"
    excel_path = tmp_path / "converted.xlsx"

    extracted = pdf_service.extract_pdf_content(pdf_path)
    assert extracted[0]["tables"] == [[
        ["Item", "Qty", "Price"],
        ["Apple", "2", "$3"],
        ["Orange", "4", "$6"],
    ]]
    assert extracted[0]["text"] == "Invoice"

    word_converter.convert_pdf_to_word(pdf_path, str(word_path))
    powerpoint_converter.convert_pdf_to_ppt(pdf_path, str(powerpoint_path))
    excel_converter.convert_pdf_to_excel(pdf_path, str(excel_path))

    word_document = Document(word_path)
    assert word_document.tables == []
    assert word_document.sections[0].header._element.xpath(".//wp:anchor")
    assert word_document.sections[0].header._element.xpath(".//wp:anchor/wp:wrapNone")
    assert not [paragraph.text for paragraph in word_document.paragraphs if paragraph.text]
    with ZipFile(word_path) as package:
        assert page_image in [package.read(name) for name in package.namelist() if name.startswith("word/media/")]

    presentation = Presentation(powerpoint_path)
    slide_shapes = presentation.slides[0].shapes
    assert not any(shape.has_table for shape in slide_shapes)
    assert len(slide_shapes) == 1
    assert slide_shapes[0].image.blob == page_image

    worksheet = load_workbook(excel_path, data_only=True)["P1_T1"]
    assert [[cell.value for cell in row] for row in worksheet.iter_rows(min_row=1, max_row=3, max_col=3)] == [
        ["Item", "Qty", "Price"],
        ["Apple", "2", "$3"],
        ["Orange", "4", "$6"],
    ]
    notes = load_workbook(excel_path, data_only=True)["Notes"]
    assert notes["A2"].value == "Invoice"


def test_powerpoint_creates_one_image_slide_per_pdf_page(tmp_path, monkeypatch):
    page_images = []
    for size, color in (((612, 792), "white"), ((792, 612), "gray")):
        image = Image.new("RGB", size, color)
        output = BytesIO()
        image.save(output, format="PNG")
        page_images.append(output.getvalue())
    monkeypatch.setattr(
        powerpoint_converter,
        "render_pdf_pages",
        lambda _: page_images,
    )
    output_path = tmp_path / "pages.pptx"

    powerpoint_converter.convert_pdf_to_ppt("source.pdf", str(output_path))

    presentation = Presentation(output_path)
    assert len(presentation.slides) == len(page_images)
    for slide, image_bytes in zip(presentation.slides, page_images):
        assert len(slide.shapes) == 1
        assert slide.shapes[0].image.blob == image_bytes


def test_symbol_font_bullet_is_converted_to_unicode_bullet():
    word = pdf_service.normalize_pdf_word(
        {"text": "\uf0b7", "fontname": "Symbol", "size": 10.5}
    )

    assert word["text"] == "\u2022"
    assert word["fontname"] == "Symbol"
