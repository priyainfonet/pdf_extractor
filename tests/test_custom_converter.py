import io
import json
import uuid

from openpyxl import load_workbook

from app.services import custom_pdf_excel
from tests.helpers import auth_headers, build_minimal_pdf, register_and_login


class FakeRow:
    def __init__(self, cells):
        self.cells = cells


class FakeTable:
    def __init__(self):
        self.rows = [
            FakeRow([(0, 10, 33, 20), (33, 10, 66, 20), (66, 10, 100, 20)]),
            FakeRow([(0, 20, 33, 30), (33, 20, 66, 30), (66, 20, 100, 30)]),
            FakeRow([(0, 30, 33, 40), (33, 30, 66, 40), (66, 30, 100, 50)]),
            FakeRow([(0, 40, 33, 50), (33, 40, 66, 50), (66, 50, 100, 60)]),
        ]

    def extract(self):
        return [
            ["Invoice No", "Description", "Quantity"],
            ["INV-1", "Widget", "10"],
            ["INV-2", "Gadget", "20"],
            ["INV-3", "Cable", "15"],
        ]


class FakePage:
    width = 100
    height = 100

    def find_tables(self, table_settings):
        return [FakeTable()]

    def extract_words(self, **kwargs):
        return []


class FakePdf:
    pages = [FakePage()]

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_value, traceback):
        return False


def _mapping(text, header, column, x0, x1):
    return {
        "pageNumber": 1,
        "text": text,
        "header": header,
        "column": column,
        "bbox": {"x0": x0, "y0": 0.1, "x1": x1, "y1": 0.2},
    }


def test_custom_export_places_extracted_table_columns(client, monkeypatch):
    monkeypatch.setattr(custom_pdf_excel.pdfplumber, "open", lambda _: FakePdf())
    access_token, _, _ = register_and_login(client, f"custom-{uuid.uuid4().hex[:10]}@example.com")
    mappings = [
        _mapping("Invoice No", "Invoice No", "A", 0.01, 0.32),
        _mapping("Description", "Description", "B", 0.34, 0.64),
        _mapping("Quantity", "Quantity", "C", 0.67, 0.95),
    ]

    response = client.post(
        "/api/custom-converter/export",
        files={"file": ("source.pdf", build_minimal_pdf(), "application/pdf")},
        data={"mappings": json.dumps(mappings)},
        headers=auth_headers(access_token),
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(response.content), data_only=True)
    worksheet = workbook["Extracted data"]
    assert [worksheet.cell(1, column).value for column in range(1, 4)] == [
        "Invoice No",
        "Description",
        "Quantity",
    ]
    assert [worksheet.cell(row, 3).value for row in range(2, 5)] == ["10", "20", "15"]


def test_custom_export_requires_authentication(client):
    response = client.post(
        "/api/custom-converter/export",
        files={"file": ("source.pdf", build_minimal_pdf(), "application/pdf")},
        data={"mappings": "[]"},
    )

    assert response.status_code == 401


def test_custom_export_rejects_duplicate_columns(client):
    access_token, _, _ = register_and_login(client, f"custom-{uuid.uuid4().hex[:10]}@example.com")
    mappings = [
        _mapping("Invoice No", "Invoice No", "A", 0.01, 0.32),
        _mapping("Quantity", "Quantity", "A", 0.67, 0.95),
    ]

    response = client.post(
        "/api/custom-converter/export",
        files={"file": ("source.pdf", build_minimal_pdf(), "application/pdf")},
        data={"mappings": json.dumps(mappings)},
        headers=auth_headers(access_token),
    )

    assert response.status_code == 422
    assert "different Excel column" in response.json()["detail"]


def test_custom_export_uses_ocr_words_for_scanned_pages(client, monkeypatch):
    monkeypatch.setattr(custom_pdf_excel.pdfplumber, "open", lambda _: FakePdf())
    monkeypatch.setattr(FakePage, "find_tables", lambda self, table_settings: [])
    access_token, _, _ = register_and_login(client, f"custom-{uuid.uuid4().hex[:10]}@example.com")
    mappings = [_mapping("Quantity", "Quantity", "C", 0.68, 0.86)]
    ocr_pages = [
        {
            "pageNumber": 1,
            "width": 1000,
            "height": 1000,
            "words": [
                {"text": "Invoice", "bbox": {"x0": 100, "y0": 100, "x1": 220, "y1": 130}},
                {"text": "Description", "bbox": {"x0": 300, "y0": 100, "x1": 550, "y1": 130}},
                {"text": "Quantity", "bbox": {"x0": 700, "y0": 100, "x1": 820, "y1": 130}},
                {"text": "INV-1", "bbox": {"x0": 100, "y0": 200, "x1": 220, "y1": 230}},
                {"text": "Widget", "bbox": {"x0": 300, "y0": 200, "x1": 500, "y1": 230}},
                {"text": "10", "bbox": {"x0": 720, "y0": 200, "x1": 750, "y1": 230}},
                {"text": "INV-2", "bbox": {"x0": 100, "y0": 300, "x1": 220, "y1": 330}},
                {"text": "Gadget", "bbox": {"x0": 300, "y0": 300, "x1": 500, "y1": 330}},
                {"text": "20", "bbox": {"x0": 720, "y0": 300, "x1": 750, "y1": 330}},
            ],
        }
    ]

    response = client.post(
        "/api/custom-converter/export",
        files={"file": ("scanned.pdf", build_minimal_pdf(), "application/pdf")},
        data={"mappings": json.dumps(mappings), "ocr_pages": json.dumps(ocr_pages)},
        headers=auth_headers(access_token),
    )

    assert response.status_code == 200, response.text
    worksheet = load_workbook(io.BytesIO(response.content), data_only=True)["Extracted data"]
    assert [worksheet.cell(row, 3).value for row in range(2, 4)] == ["10", "20"]


def test_custom_export_extracts_only_nearby_gstin(client, monkeypatch):
    monkeypatch.setattr(custom_pdf_excel.pdfplumber, "open", lambda _: FakePdf())
    monkeypatch.setattr(FakePage, "find_tables", lambda self, table_settings: [])
    access_token, _, _ = register_and_login(client, f"custom-{uuid.uuid4().hex[:10]}@example.com")
    mappings = [_mapping("GSTIN", "GSTIN", "A", 0.08, 0.24)]
    ocr_pages = [
        {
            "pageNumber": 1,
            "width": 1000,
            "height": 1000,
            "words": [
                {"text": "GSTIN", "bbox": {"x0": 100, "y0": 100, "x1": 200, "y1": 130}},
                {"text": "29ABCDE1234F1Z5", "bbox": {"x0": 240, "y0": 100, "x1": 450, "y1": 130}},
                {"text": "State:", "bbox": {"x0": 100, "y0": 180, "x1": 160, "y1": 210}},
                {"text": "Karnataka", "bbox": {"x0": 170, "y0": 180, "x1": 300, "y1": 210}},
                {"text": "Invoice", "bbox": {"x0": 100, "y0": 240, "x1": 190, "y1": 270}},
                {"text": "No:", "bbox": {"x0": 200, "y0": 240, "x1": 240, "y1": 270}},
                {"text": "S-191/2022-23", "bbox": {"x0": 250, "y0": 240, "x1": 430, "y1": 270}},
                {"text": "HSN/SAC", "bbox": {"x0": 100, "y0": 300, "x1": 220, "y1": 330}},
                {"text": "998596", "bbox": {"x0": 230, "y0": 300, "x1": 300, "y1": 330}},
            ],
        }
    ]

    response = client.post(
        "/api/custom-converter/export",
        files={"file": ("invoice.pdf", build_minimal_pdf(), "application/pdf")},
        data={"mappings": json.dumps(mappings), "ocr_pages": json.dumps(ocr_pages)},
        headers=auth_headers(access_token),
    )

    assert response.status_code == 200, response.text
    worksheet = load_workbook(io.BytesIO(response.content), data_only=True)["Extracted data"]
    assert worksheet["A2"].value == "29ABCDE1234F1Z5"
    assert worksheet["A3"].value is None