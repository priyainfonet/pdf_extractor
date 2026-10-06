"""Converts extracted PDF content into a formatted .xlsx workbook."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter, column_index_from_string

from app.services.pdf_service import (
    extract_pdf_content,
    find_primary_table,
    normalize_headers,
    drop_blank_rows,
    build_content_blocks,
)

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

# Below this fraction of non-empty cells, a table is treated as a dense
# bordered FORM (many small label/value boxes merged into one giant
# sparse grid by pdfplumber) rather than a normal well-formed table, and
# gets decomposed into real sub-tables + plain rows instead of dumped
# as one huge mostly-blank sheet. Verified against a real customs form
# (12.5% fill) vs. a normal invoice table with one blank optional cell
# (91.7% fill) — comfortable separation.
FORM_LIKE_FILL_THRESHOLD = 0.40


def _fill_ratio(table: list[list[str]]) -> float:
    total = sum(len(row) for row in table)
    if not total:
        return 1.0
    non_empty = sum(1 for row in table for c in row if c and c.strip())
    return non_empty / total


def _write_simple_table(ws, table: list[list[str]]) -> int:
    """The original behavior — used for well-formed tables. Returns the
    number of columns used, for width-sizing."""
    for row_idx, row in enumerate(table, start=1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if row_idx == 1:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    return max((len(row) for row in table), default=0)


def _write_form_blocks(ws, table: list[list[str]]) -> int:
    """For dense bordered forms: real header+data sub-tables get their
    own mini-table styling; everything else is shown as a plain row of
    values with no guessed label/value relationship (see
    build_content_blocks's docstring for why that guess was dropped —
    it produced confidently wrong pairings on real test documents)."""
    blocks = build_content_blocks(table)
    current_row = 1
    max_col = 0

    for block_type, first, rest in blocks:
        if block_type == "table":
            header, data_rows = first, rest
            for col_idx, value in enumerate(header, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=value)
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center")
            max_col = max(max_col, len(header))
            current_row += 1

            for data_row in data_rows:
                for col_idx, value in enumerate(data_row, start=1):
                    ws.cell(row=current_row, column=col_idx, value=value)
                max_col = max(max_col, len(data_row))
                current_row += 1
        else:  # "row" — plain values, no assumed structure
            values = first
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row=current_row, column=col_idx, value=value)
            max_col = max(max_col, len(values))
            current_row += 1

        current_row += 1  # blank separator row between blocks

    return max_col


def convert_pdf_to_excel(pdf_path: str, output_path: str) -> str:
    pages_data = extract_pdf_content(pdf_path)

    wb = Workbook()
    wb.remove(wb.active)

    notes_lines = []
    sheet_count = 0

    for page in pages_data:
        if page["text"].strip():
            notes_lines.append(f"--- Page {page['page_number']} ---")
            notes_lines.append(page["text"])

        for t_idx, table in enumerate(page["tables"], start=1):
            table = drop_blank_rows(table)
            if not table:
                continue

            sheet_count += 1
            sheet_name = f"P{page['page_number']}_T{t_idx}"[:31]
            ws = wb.create_sheet(title=sheet_name)

            if _fill_ratio(table) < FORM_LIKE_FILL_THRESHOLD:
                max_col = _write_form_blocks(ws, table)
                last_row = ws.max_row
            else:
                max_col = _write_simple_table(ws, table)
                last_row = len(table)

            for col_idx in range(1, max_col + 1):
                col_letter = get_column_letter(col_idx)
                max_len = max(
                    (len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, last_row + 1)),
                    default=10,
                )
                ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

    if notes_lines:
        notes_ws = wb.create_sheet(title="Notes")
        for i, line in enumerate(notes_lines, start=1):
            notes_ws.cell(row=i, column=1, value=line)
        notes_ws.column_dimensions["A"].width = 100

    if sheet_count == 0 and not notes_lines:
        wb.create_sheet(title="Sheet1").cell(row=1, column=1, value="No tables or text found in PDF.")

    wb.save(output_path)
    return output_path


def convert_pdf_to_excel_with_mapping(pdf_path: str, output_path: str, mapping: dict[str, str]) -> str:
    """Like convert_pdf_to_excel, but places the primary table's columns
    at whichever Excel column letter the user chose in the mapping UI,
    instead of the default left-to-right one-sheet-per-table layout.

    mapping is {pdf_heading_name: excel_column_letter}. Each heading's
    data stays tied to it by its ORIGINAL column index in the extracted
    table — only its destination column changes, never which values
    belong to it.
    """
    pages_data = extract_pdf_content(pdf_path)
    table = find_primary_table(pages_data)

    if not table:
        raise ValueError("No table found in this PDF to apply a column mapping to.")

    headers = normalize_headers(table[0])
    data_rows = drop_blank_rows(table[1:])

    wb = Workbook()
    ws = wb.active
    ws.title = "Mapped Data"

    used_columns = []
    for heading_index, heading_name in enumerate(headers):
        target_letter = mapping.get(heading_name)
        if not target_letter:
            # Heading wasn't included in the mapping the frontend sent —
            # skip it rather than guess a placement.
            continue

        col_idx = column_index_from_string(target_letter)
        used_columns.append(col_idx)

        header_cell = ws.cell(row=1, column=col_idx, value=heading_name)
        header_cell.font = HEADER_FONT
        header_cell.fill = HEADER_FILL
        header_cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_offset, row in enumerate(data_rows, start=2):
            value = row[heading_index] if heading_index < len(row) else ""
            ws.cell(row=row_offset, column=col_idx, value=value)

    for col_idx in used_columns:
        letter = get_column_letter(col_idx)
        max_len = max(
            (len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, len(data_rows) + 2)),
            default=10,
        )
        ws.column_dimensions[letter].width = min(max_len + 4, 50)

    ws.freeze_panes = "A2"

    wb.save(output_path)
    return output_path
