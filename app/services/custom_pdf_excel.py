"""Coordinate-aware extraction for the custom PDF-to-Excel workflow."""

import io
import re
from difflib import SequenceMatcher

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter


HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
GSTIN_PATTERN = re.compile(r"(?<![A-Z0-9])(\d{2}[A-Z]{5}\d{4}[A-Z][A-Z0-9]Z[A-Z0-9])(?![A-Z0-9])", re.IGNORECASE)


class CustomExtractionError(ValueError):
    pass


def _compact(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", text.lower())


def _box_to_points(box: dict[str, float], page) -> tuple[float, float, float, float]:
    return (
        box["x0"] * page.width,
        box["y0"] * page.height,
        box["x1"] * page.width,
        box["y1"] * page.height,
    )


def _table_candidates(page):
    for settings in (
        {"vertical_strategy": "lines", "horizontal_strategy": "lines"},
        {"vertical_strategy": "text", "horizontal_strategy": "text"},
    ):
        try:
            tables = page.find_tables(table_settings=settings)
        except Exception:
            continue
        if tables:
            return tables
    return []


def _cell_overlap_score(selection_box, cell_box) -> float:
    sx0, sy0, sx1, sy1 = selection_box
    cx0, cy0, cx1, cy1 = cell_box
    intersection_width = max(0.0, min(sx1, cx1) - max(sx0, cx0))
    intersection_height = max(0.0, min(sy1, cy1) - max(sy0, cy0))
    intersection = intersection_width * intersection_height
    selection_area = max(1e-9, (sx1 - sx0) * (sy1 - sy0))
    center_x = (sx0 + sx1) / 2
    center_y = (sy0 + sy1) / 2
    center_inside = cx0 <= center_x <= cx1 and cy0 <= center_y <= cy1
    return intersection / selection_area + (1.0 if center_inside else 0.0)


def _matching_table_column(page, selection, selection_box):
    title = _compact(selection.text)
    best = None

    for table in _table_candidates(page):
        extracted = table.extract() or []
        for row_index, row in enumerate(table.rows):
            for column_index, cell_box in enumerate(row.cells):
                if cell_box is None or row_index >= len(extracted):
                    continue
                cell_text = extracted[row_index][column_index] if column_index < len(extracted[row_index]) else ""
                cell_compact = _compact(cell_text or "")
                geometry_score = _cell_overlap_score(selection_box, cell_box)
                if geometry_score <= 0:
                    continue

                text_score = SequenceMatcher(None, title, cell_compact).ratio() if title and cell_compact else 0
                if title and cell_compact and (title in cell_compact or cell_compact in title):
                    text_score = 1.0
                if text_score < 0.65:
                    continue
                score = geometry_score + text_score * 2
                if best is None or score > best[0]:
                    best = (score, extracted, row_index, column_index)

    if best is None:
        return None

    _, rows, header_row, column = best
    values = []
    for row in rows[header_row + 1 :]:
        value = row[column].strip() if column < len(row) and row[column] else ""
        if value and _compact(value) != title:
            values.append(value)
    return values


def _group_lines(words: list[dict]) -> list[list[dict]]:
    lines: list[list[dict]] = []
    for word in sorted(words, key=lambda item: (item["top"], item["x0"])):
        word_center = (word["top"] + word["bottom"]) / 2
        matching_line = next(
            (
                line
                for line in lines
                if abs(word_center - sum((item["top"] + item["bottom"]) / 2 for item in line) / len(line))
                <= max(3.0, (word["bottom"] - word["top"]) * 0.6)
            ),
            None,
        )
        if matching_line is None:
            lines.append([word])
        else:
            matching_line.append(word)
    return [sorted(line, key=lambda item: item["x0"]) for line in lines]


def _group_header_words(words: list[dict]) -> list[list[dict]]:
    if not words:
        return []
    heights = [max(1.0, word["bottom"] - word["top"]) for word in words]
    gap_limit = max(6.0, sorted(heights)[len(heights) // 2] * 1.5)
    groups: list[list[dict]] = [[words[0]]]
    for word in words[1:]:
        previous = groups[-1][-1]
        if word["x0"] - previous["x1"] > gap_limit:
            groups.append([word])
        else:
            groups[-1].append(word)
    return groups


def _page_words(page, ocr_page=None) -> list[dict]:
    if ocr_page:
        scale_x = page.width / ocr_page.width
        scale_y = page.height / ocr_page.height
        return [
            {
                "text": word.text,
                "x0": word.bbox.x0 * scale_x,
                "x1": word.bbox.x1 * scale_x,
                "top": word.bbox.y0 * scale_y,
                "bottom": word.bbox.y1 * scale_y,
            }
            for word in ocr_page.words
            if word.text.strip()
        ]
    return page.extract_words(x_tolerance=2, y_tolerance=3) or []


def _extract_gstin(page, selection, selection_box, ocr_page=None) -> str | None:
    label = _compact(selection.text)
    if not any(term in label for term in ("gstin", "gstno", "gstnumber", "goodsandservicestaxidentificationnumber")):
        return None

    sx0, sy0, sx1, sy1 = selection_box
    candidates = []
    for line in _group_lines(_page_words(page, ocr_page)):
        line_text = " ".join(word["text"] for word in line)
        match = GSTIN_PATTERN.search(line_text)
        if not match:
            continue

        line_top = min(word["top"] for word in line)
        line_bottom = max(word["bottom"] for word in line)
        if line_bottom < sy0 - 2:
            continue
        if line_top > sy1 + max(72, page.height * 0.12):
            continue

        same_line = line_top <= sy1 + 2 and line_bottom >= sy0 - 2
        if same_line:
            line_left = min(word["x0"] for word in line)
            line_right = max(word["x1"] for word in line)
            horizontal_distance = max(0, line_left - sx1, sx0 - line_right)
            score = horizontal_distance
        else:
            score = 1000 + max(0, line_top - sy1)
        candidates.append((score, match.group(1).upper()))

    return min(candidates, default=(0, None), key=lambda candidate: candidate[0])[1]


def _word_values(page, selection, selection_box, ocr_page=None) -> list[str]:
    words = _page_words(page, ocr_page)

    lines = _group_lines(words)
    title = _compact(selection.text)
    sx0, sy0, sx1, sy1 = selection_box
    selection_center_x = (sx0 + sx1) / 2
    selection_center_y = (sy0 + sy1) / 2

    header_line = None
    header_groups = None
    best_score = -1.0
    for line in lines:
        line_top = min(word["top"] for word in line)
        line_bottom = max(word["bottom"] for word in line)
        if not (line_top - 4 <= selection_center_y <= line_bottom + 4):
            continue
        groups = _group_header_words(line)
        for group in groups:
            group_text = " ".join(word["text"] for word in group)
            group_compact = _compact(group_text)
            text_score = SequenceMatcher(None, title, group_compact).ratio() if title and group_compact else 0
            if title and group_compact and (title in group_compact or group_compact in title):
                text_score = 1.0
            if text_score < 0.65:
                continue
            group_center_x = (group[0]["x0"] + group[-1]["x1"]) / 2
            geometry_score = _cell_overlap_score(
                selection_box,
                (group[0]["x0"], line_top, group[-1]["x1"], line_bottom),
            )
            score = text_score * 3 + geometry_score - abs(group_center_x - selection_center_x) / max(1, page.width)
            if score > best_score:
                best_score = score
                header_line = line
                header_groups = groups

    if not header_line or not header_groups or best_score < 1:
        return []

    selected_group_index = min(
        range(len(header_groups)),
        key=lambda index: abs(
            (header_groups[index][0]["x0"] + header_groups[index][-1]["x1"]) / 2 - selection_center_x
        ),
    )
    selected_group = header_groups[selected_group_index]
    selected_center = (selected_group[0]["x0"] + selected_group[-1]["x1"]) / 2
    left = 0.0
    right = page.width
    if selected_group_index:
        previous = header_groups[selected_group_index - 1]
        previous_center = (previous[0]["x0"] + previous[-1]["x1"]) / 2
        left = (previous_center + selected_center) / 2
    if selected_group_index + 1 < len(header_groups):
        following = header_groups[selected_group_index + 1]
        following_center = (following[0]["x0"] + following[-1]["x1"]) / 2
        right = (selected_center + following_center) / 2

    header_bottom = max(word["bottom"] for word in header_line)
    values = []
    for line in lines:
        if min(word["top"] for word in line) <= header_bottom + 1:
            continue
        selected_words = [
            word
            for word in line
            if left <= (word["x0"] + word["x1"]) / 2 < right
        ]
        value = " ".join(word["text"] for word in selected_words).strip()
        if value:
            values.append(value)
    return values


def _extract_selection(pdf, selection, ocr_pages_by_number) -> list[str]:
    if selection.page_number > len(pdf.pages):
        raise CustomExtractionError(f"Page {selection.page_number} is not present in the uploaded PDF.")

    page = pdf.pages[selection.page_number - 1]
    selection_box = _box_to_points(selection.bbox.model_dump(), page)
    gstin = _extract_gstin(page, selection, selection_box, ocr_pages_by_number.get(selection.page_number))
    if gstin:
        return [gstin]

    table_values = _matching_table_column(page, selection, selection_box)
    if table_values:
        return table_values

    ocr_page = ocr_pages_by_number.get(selection.page_number)
    values = _word_values(page, selection, selection_box, ocr_page)
    if values:
        return values

    raise CustomExtractionError(
        f"No data was found below '{selection.text}' on page {selection.page_number}. "
        "Check the selected heading and its bounding box."
    )


def build_custom_workbook(pdf_bytes: bytes, selections: list, ocr_pages: list) -> bytes:
    """Extract values below mapped headings and return an in-memory XLSX."""
    ocr_pages_by_number = {page.page_number: page for page in ocr_pages}
    try:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            values_by_column = [
                (selection, _extract_selection(pdf, selection, ocr_pages_by_number))
                for selection in selections
            ]
    except CustomExtractionError:
        raise
    except Exception as exc:
        raise CustomExtractionError("The uploaded PDF could not be read for extraction.") from exc

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Extracted data"

    max_row_count = max(len(values) for _, values in values_by_column)
    for selection, values in values_by_column:
        column_index = column_index_from_string(selection.column)
        header = worksheet.cell(row=1, column=column_index, value=selection.header)
        header.font = HEADER_FONT
        header.fill = HEADER_FILL
        header.alignment = Alignment(horizontal="center", vertical="center")
        for row_index in range(max_row_count):
            value = values[row_index] if row_index < len(values) else None
            worksheet.cell(row=row_index + 2, column=column_index, value=value)
        letter = get_column_letter(column_index)
        longest_value = max([len(selection.header), *(len(value) for value in values)], default=10)
        worksheet.column_dimensions[letter].width = min(longest_value + 4, 50)

    worksheet.freeze_panes = "A2"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()