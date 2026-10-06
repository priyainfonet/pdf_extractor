"""Converts PDF pages into visually faithful PowerPoint slides."""

from io import BytesIO

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.services.pdf_service import render_pdf_pages

SLIDE_WIDTH = Inches(13.33)


def convert_pdf_to_ppt(pdf_path: str, output_path: str) -> str:
    page_images = render_pdf_pages(pdf_path)

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    if page_images:
        with Image.open(BytesIO(page_images[0])) as first_page:
            prs.slide_height = int(SLIDE_WIDTH * first_page.height / first_page.width)
    blank_layout = prs.slide_layouts[6]

    for page_image in page_images:
        slide = prs.slides.add_slide(blank_layout)
        with Image.open(BytesIO(page_image)) as image:
            scale = min(prs.slide_width / image.width, prs.slide_height / image.height)
            width = int(image.width * scale)
            height = int(image.height * scale)
        left = (prs.slide_width - width) // 2
        top = (prs.slide_height - height) // 2
        slide.shapes.add_picture(BytesIO(page_image), left, top, width=width, height=height)

    prs.save(output_path)
    return output_path
