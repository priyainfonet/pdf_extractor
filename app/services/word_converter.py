"""Converts PDF pages into visually faithful Word document pages."""

from copy import deepcopy
from io import BytesIO

from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from app.services.pdf_service import render_pdf_pages


def _configure_section(section, width: float, height: float) -> None:
    section.page_width = Pt(width)
    section.page_height = Pt(height)
    section.top_margin = Pt(0)
    section.bottom_margin = Pt(0)
    section.left_margin = Pt(0)
    section.right_margin = Pt(0)
    section.header_distance = Pt(0)
    section.footer_distance = Pt(0)


def _add_graphics_background(section, image_bytes: bytes) -> None:
    section.header.is_linked_to_previous = False
    paragraph = section.header.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    picture = paragraph.add_run().add_picture(
        BytesIO(image_bytes), width=section.page_width, height=section.page_height
    )
    inline = picture._inline
    anchor = OxmlElement("wp:anchor")
    for name, value in (
        ("distT", "0"),
        ("distB", "0"),
        ("distL", "0"),
        ("distR", "0"),
        ("simplePos", "0"),
        ("relativeHeight", "0"),
        ("behindDoc", "1"),
        ("locked", "0"),
        ("layoutInCell", "1"),
        ("allowOverlap", "1"),
    ):
        anchor.set(name, value)

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)
    for tag in ("wp:positionH", "wp:positionV"):
        position = OxmlElement(tag)
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = "0"
        position.append(offset)
        anchor.append(position)

    for tag in (
        "wp:extent",
        "wp:effectExtent",
        "wp:wrapNone",
        "wp:docPr",
        "wp:cNvGraphicFramePr",
        "a:graphic",
    ):
        element = inline.find(qn(tag))
        if element is not None:
            anchor.append(deepcopy(element))
        elif tag == "wp:wrapNone":
            anchor.append(OxmlElement(tag))

    inline.getparent().replace(inline, anchor)


def convert_pdf_to_word(pdf_path: str, output_path: str) -> str:
    page_images = render_pdf_pages(pdf_path)

    doc = Document()
    for index, image_bytes in enumerate(page_images):
        with Image.open(BytesIO(image_bytes)) as image:
            page_width = image.width / 2
            page_height = image.height / 2
        if index:
            section = doc.add_section(WD_SECTION.NEW_PAGE)
            _configure_section(section, page_width, page_height)
        else:
            _configure_section(doc.sections[0], page_width, page_height)
        _add_graphics_background(section if index else doc.sections[0], image_bytes)

    doc.save(output_path)
    return output_path
